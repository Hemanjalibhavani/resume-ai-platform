"""PDF/DOCX -> text -> clean text -> section detection -> information extraction."""
import io
import re

from backend.services.skill_extractor import extract_skills
from backend.utils.helpers import AppError, get_logger
from backend.utils.text_cleaner import clean_text

log = get_logger(__name__)

SECTION_ALIASES = {
    "summary": ["summary", "professional summary", "profile", "about me", "career summary"],
    "objective": ["objective", "career objective"],
    "education": ["education", "academic background", "academics", "qualifications", "educational qualifications"],
    "skills": ["skills", "technical skills", "key skills", "core competencies", "technologies", "tech stack"],
    "projects": ["projects", "academic projects", "personal projects", "project work", "key projects"],
    "internships": ["internships", "internship", "internship experience", "industrial training", "training"],
    "experience": ["experience", "work experience", "professional experience", "employment history"],
    "certifications": ["certifications", "certificates", "licenses", "courses", "certifications and courses"],
    "achievements": ["achievements", "awards", "honors", "accomplishments", "awards and achievements"],
    "extracurricular": ["extracurricular", "extra curricular", "extracurricular activities", "extracurricular hackathon",
                        "extracurricular hackathons", "hackathon", "hackathons", "activities",
                        "volunteering", "positions of responsibility"],
    "languages": ["languages", "language proficiency"],
    "hobbies": ["hobbies", "interests", "hobbies and interests"],
}
_ALIAS_TO_SECTION = {a: s for s, al in SECTION_ALIASES.items() for a in al}
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(\+?\d[\d\s\-()]{8,16}\d)")
BULLET_RE = re.compile(r"^\s*[-*]\s+")


def extract_text(data: bytes, filename: str) -> str:
    """Extract raw text from PDF, DOCX or TXT bytes."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if ext == "pdf":
            try:
                import pymupdf as fitz  # PyMuPDF >= 1.24
            except ImportError:  # older releases
                import fitz

            with fitz.open(stream=data, filetype="pdf") as doc:
                text = "\n".join(page.get_text("text") for page in doc)
        elif ext == "docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif ext == "txt":
            text = data.decode("utf-8", errors="ignore")
        else:
            raise AppError("Unsupported file type.", 400)
    except AppError:
        raise
    except Exception as exc:  # corrupted / encrypted files
        log.warning("Text extraction failed: %s", exc)
        raise AppError("Please upload a valid PDF resume." if ext == "pdf"
                       else "Please upload a valid DOCX file.", 400) from exc
    text = clean_text(text)
    if len(text) < 30:
        raise AppError("We could not extract readable text from this file.", 422)
    return text


def _heading_of(line: str):
    """Return (section, inline_rest) if the line starts a known section."""
    stripped = line.strip()
    if not stripped or len(stripped) > 60:
        return None
    key = re.sub(r"\s+", " ", re.sub(r"[^a-z ]", "", stripped.lower().split(":")[0])).strip()
    if key in _ALIAS_TO_SECTION:
        rest = stripped.split(":", 1)[1].strip() if ":" in stripped else ""
        return _ALIAS_TO_SECTION[key], rest
    return None


def detect_sections(text: str) -> dict[str, str]:
    """Split resume text into named sections; the lines before any heading are 'contact'."""
    sections: dict[str, list[str]] = {"contact": []}
    current = "contact"
    for line in text.split("\n"):
        head = _heading_of(line)
        if head:
            current = head[0]
            sections.setdefault(current, [])
            if head[1]:
                sections[current].append(head[1])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


SENTENCE_END = (".", "!", "?")
EDU_WORDS = re.compile(r"college|school|university|institute|academy|polytechnic", re.I)


def _entries(content: str) -> list[dict]:
    """Group lines into entries.

    Non-bullet line = entry title; short follow-up lines (dates, stack, 'Live Demo') = meta;
    bullet lines = description; wrapped PDF lines are joined to the bullet they continue.
    """
    entries: list[dict] = []
    in_bullet = False
    for raw in [l for l in content.split("\n") if l.strip()]:
        line = raw.strip()
        if BULLET_RE.match(line):
            body = BULLET_RE.sub("", line).strip()
            if entries:
                entries[-1]["description"] = (entries[-1]["description"] + " " + body).strip()
            else:
                entries.append({"title": body[:80], "description": body, "meta": ""})
            in_bullet = True
            continue
        if entries and in_bullet and (line[0].islower() or not entries[-1]["description"].endswith(SENTENCE_END)):
            entries[-1]["description"] += " " + line          # wrapped continuation line
            continue
        if entries and not entries[-1]["description"] and len(line) <= 70 and not line.endswith("."):
            entries[-1]["meta"] = (entries[-1]["meta"] + " | " + line).strip(" |")
            continue
        entries.append({"title": line, "description": "", "meta": ""})
        in_bullet = False
    return entries


def _lines(content: str) -> list[str]:
    return [BULLET_RE.sub("", l).strip() for l in content.split("\n") if l.strip()]


def _items(content: str) -> list[str]:
    """Certification-style lists: items separated by '|' (possibly wrapped over lines) or one per line."""
    if "|" in content:
        flat = re.sub(r"\s*\n\s*", " ", content)
        return [i.strip(" -") for i in flat.split("|") if i.strip(" -")]
    return _lines(content)


def _education(content: str) -> list[str]:
    """Merge institution / years / degree / CGPA lines into one entry per institution."""
    lines = _lines(content)
    if not any(EDU_WORDS.search(l) for l in lines):
        return lines
    out: list[str] = []
    for l in lines:
        if EDU_WORDS.search(l) or not out:
            out.append(l)
        else:
            out[-1] += " | " + l
    return out


def extract_contact(text: str, head: str) -> dict:
    email = EMAIL_RE.search(text)
    phone = ""
    for m in PHONE_RE.finditer(head or text):
        digits = re.sub(r"\D", "", m.group(1))
        if 10 <= len(digits) <= 13:
            phone = m.group(1).strip()
            break
    name = "Unknown"
    for line in [l.strip() for l in (head or text).split("\n") if l.strip()][:6]:
        words = line.split()
        if ("@" in line or re.search(r"\d", line) or not 1 < len(words) <= 4
                or re.search(r"resume|curriculum|vitae|http", line, re.I)):
            continue
        if all(re.fullmatch(r"[A-Za-z.'-]+", w) for w in words):
            name = line.title() if line.isupper() else line
            break
    link = lambda pat: (re.search(pat, text, re.I) or [None])[0]
    linkedin = link(r"linkedin\.com/[\w/-]+") or ("LinkedIn (link in resume)" if re.search(r"\blinkedin\b", text, re.I) else "")
    github = link(r"github\.com/[\w-]+") or ("GitHub (link in resume)" if re.search(r"\bgithub\b", text, re.I) else "")
    return {"name": name, "email": email.group(0) if email else "", "phone": phone,
            "linkedin": linkedin, "github": github}


def degree_level(text: str) -> int:
    """0 unknown, 1 diploma, 2 bachelor, 3 master, 4 doctorate."""
    t = text.lower()
    if re.search(r"ph\.?d|doctorate", t):
        return 4
    if re.search(r"\bm\.?\s?tech\b|master|\bm\.?sc\b|\bmba\b|\bmca\b|\bm\.?e\b", t):
        return 3
    if re.search(r"\bb\.?\s?tech\b|bachelor|\bb\.?e\b|\bb\.?sc\b|\bbca\b|\bb\.?com\b|undergraduate", t):
        return 2
    if "diploma" in t:
        return 1
    return 0


MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_MON = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?"
RANGE_RE = re.compile(rf"\b{_MON}\s+(\d{{4}})\s*(?:-|–|—|to)\s*(?:{_MON}\s+(\d{{4}})|(present|current|ongoing))", re.I)


def months_from_ranges(text: str) -> int:
    """Sum the length of 'Mon YYYY - Mon YYYY' ranges (e.g. 'Jun 2025 - Sep 2025' = 3 months)."""
    from datetime import date

    total = 0
    for m in RANGE_RE.finditer(text):
        start = int(m.group(2)) * 12 + MONTHS[m.group(1).lower()[:3]]
        if m.group(5):
            today = date.today()
            end = today.year * 12 + today.month
        else:
            end = int(m.group(4)) * 12 + MONTHS[m.group(3).lower()[:3]]
        total += max(0, end - start)
    return total


def years_of_experience(text: str, internships: str = "") -> float:
    """Declared years + internship months (counted as full months of experience)."""
    years = 0.0
    for m in re.finditer(r"(\d+(?:\.\d+)?)\s*\+?\s*years?", text, re.I):
        ctx = text[max(0, m.start() - 40): m.end() + 40].lower()
        if "experience" in ctx:
            years = max(years, float(m.group(1)))
    months = sum(int(m.group(1)) for m in re.finditer(r"(\d+)\s*months?", internships, re.I)) + months_from_ranges(internships)
    return round(years + months / 12.0, 2)


def parse_resume(text: str) -> dict:
    """Full structured parse of cleaned resume text."""
    sections = detect_sections(text)
    contact = extract_contact(text, sections.get("contact", ""))
    skills = extract_skills(text)
    return {
        "contact": contact,
        "sections": sections,
        "sections_detected": sorted(sections.keys()),
        "education": _education(sections.get("education", "")),
        "education_level": degree_level(sections.get("education", "")),
        "skills": skills,
        "projects": _entries(sections.get("projects", "")),
        "internships": _entries(sections.get("internships", "")),
        "experience": _entries(sections.get("experience", "")),
        "certifications": _items(sections.get("certifications", "")),
        "achievements": [" – ".join(x for x in (e["title"], e["meta"], e["description"]) if x)
                         for e in _entries(sections.get("achievements", "") + "\n" + sections.get("extracurricular", ""))],
        "years_experience": years_of_experience(sections.get("experience", "") + " " + text,
                                                sections.get("internships", "") + "\n" + sections.get("experience", "")),
    }
